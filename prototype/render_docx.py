from __future__ import annotations

from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from prototype.report_schema import CallGradingReport, DIMENSIONS


_DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "call-grading-template.docx"


def _score_display_int(score_out_of_5: float) -> str:
    # The example PDF shows integer scores like "3/5". Round to nearest integer for display.
    return f"{int(round(score_out_of_5))}/5"


def _clean_flag_text(flag: str) -> str:
    flag = str(flag).strip()
    # Prevent double warning symbols if the LLM already included them.
    if flag.startswith("⚠"):
        return flag.lstrip("⚠").strip()
    return flag


def _join_evidence_notes(evidence: str, notes: str) -> str:
    evidence = str(evidence).strip()
    notes = str(notes).strip()
    if evidence and notes:
        return f"{evidence} {notes}"
    return evidence or notes


def _set_cell_shading(cell, *, fill_hex: str) -> None:
    """
    Set Word table cell background color.

    python-docx does not have a simple public API for cell shading, so we edit tcPr/w:shd.
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


def _score_to_color_hex(score_out_of_5: float) -> str:
    # Color coding: green for good, yellow for ok, red for poor.
    # Thresholds aligned with `rating_from_score` in `report_schema.py`.
    if score_out_of_5 >= 4.0:
        return "C8E6C9"  # green
    if score_out_of_5 >= 2.5:
        return "FFF59D"  # yellow
    return "FFCDD2"  # red


def _find_heading_paragraph_index(doc: Document, heading_text: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == heading_text:
            return i
    raise ValueError(f"Heading not found in template: '{heading_text}'")


def write_call_grading_docx(*, assessed: CallGradingReport, out_path: str, template_path: Optional[str] = None) -> None:
    template_file = Path(template_path) if template_path else _DEFAULT_TEMPLATE_PATH
    if not template_file.exists():
        raise FileNotFoundError(
            f"Missing template docx: {template_file}. "
            "Run `python prototype/template_builder.py` to generate it."
        )

    doc = Document(str(template_file))

    # Tables in the template (scenario2-only):
    # - 0: Call metadata
    # - 1: Overall score
    # - 2: Dimension scores
    # - 3: Critical flags (single cell)
    meta_table = doc.tables[0]
    overall_table = doc.tables[1]
    dim_table = doc.tables[2]
    critical_table = doc.tables[3]

    meta_table.cell(0, 1).text = assessed.metadata.call_id
    meta_table.cell(0, 3).text = assessed.metadata.date
    meta_table.cell(1, 1).text = assessed.metadata.operator
    meta_table.cell(1, 3).text = assessed.metadata.duration
    meta_table.cell(2, 1).text = assessed.metadata.call_type
    meta_table.cell(2, 3).text = assessed.metadata.assessed_by

    overall_table.cell(0, 0).text = (
        f"OVERALL SCORE\n{assessed.overall_score_out_of_5:.1f} / 5.0\n{assessed.rating_label}"
    )
    _set_cell_shading(overall_table.cell(0, 0), fill_hex=_score_to_color_hex(assessed.overall_score_out_of_5))

    dim_by_name = {ds.name: ds for ds in assessed.dimension_scores}
    if len(DIMENSIONS) != len(dim_table.rows) - 1:
        raise ValueError(
            f"Template dimension table row count mismatch. "
            f"Template has {len(dim_table.rows) - 1} dimensions but code expects {len(DIMENSIONS)}."
        )

    for row_idx, rubric in enumerate(DIMENSIONS, start=1):
        ds = dim_by_name.get(rubric.name)
        if not ds:
            raise ValueError(f"Missing dimension score for: {rubric.name}")

        dim_table.cell(row_idx, 0).text = rubric.name
        dim_table.cell(row_idx, 1).text = f"{rubric.weight_percent:g}%"
        score_cell = dim_table.cell(row_idx, 2)
        score_cell.text = _score_display_int(ds.score_out_of_5)
        _set_cell_shading(score_cell, fill_hex=_score_to_color_hex(ds.score_out_of_5))
        dim_table.cell(row_idx, 3).text = _join_evidence_notes(
            evidence=ds.evidence_and_notes.evidence,
            notes=ds.evidence_and_notes.notes,
        )

    critical_flags = [_clean_flag_text(f) for f in assessed.critical_flags]
    if critical_flags:
        critical_table.cell(0, 0).text = "CRITICAL FLAGS\n" + "\n".join(f"⚠ {f}" for f in critical_flags)
    else:
        critical_table.cell(0, 0).text = "CRITICAL FLAGS\nNone identified."

    # Replace strengths bullets (template has a fixed number of existing bullet paragraphs).
    strengths_heading_idx = _find_heading_paragraph_index(doc, "STRENGTHS")
    areas_heading_idx = _find_heading_paragraph_index(doc, "AREAS FOR IMPROVEMENT")
    coaching_heading_idx = _find_heading_paragraph_index(doc, "COACHING RECOMMENDATION")

    strength_paragraphs = [
        p for p in doc.paragraphs[strengths_heading_idx + 1 : areas_heading_idx] if p.text.strip().startswith("✓")
    ]
    for i, p in enumerate(strength_paragraphs):
        text = assessed.strengths[i] if i < len(assessed.strengths) else "None identified."
        p.text = f"✓ {text}"

    area_paragraphs = [
        p
        for p in doc.paragraphs[areas_heading_idx + 1 : coaching_heading_idx]
        if p.text.strip().startswith("▶")
    ]
    for i, p in enumerate(area_paragraphs):
        text = assessed.areas_for_improvement[i] if i < len(assessed.areas_for_improvement) else "None identified."
        p.text = f"▶ {text}"

    # Replace coaching recommendation paragraph (the first non-empty paragraph after the heading).
    for j in range(coaching_heading_idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[j].text.strip():
            doc.paragraphs[j].text = assessed.coaching_recommendation
            break

    # Replace footer if we can locate it.
    for p in doc.paragraphs:
        if p.text.strip().startswith("Generated by AI Quality Assurance System"):
            p.text = assessed.footer_disclaimer
            break

    doc.save(out_path)

